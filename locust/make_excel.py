"""Build capacity_results.xlsx from the three valid result folders."""
import csv
import glob
import json
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
RUNS = {"browse": "20260930-1928-browse", "coding": "20260930-2045-coding", "jupyter": "20260930-2110-jupyter"}
WARMUP = 90

F = "Arial"
HDR = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
BOLD = Font(name=F, bold=True)
NORM = Font(name=F)
INPUT = Font(name=F, color="0000FF")
LINK = Font(name=F, color="008000")
TITLE = Font(name=F, bold=True, size=14)
NOTE = Font(name=F, italic=True, color="595959", size=9)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")


def report(track):
    return json.load(open(os.path.join(HERE, "results", RUNS[track], "stages_report.json")))


def stage_windows(track):
    d = os.path.join(HERE, "results", RUNS[track])
    st = [(float(t), int(u)) for t, u in csv.reader(open(os.path.join(d, "stages.csv")))]
    reqs = [r for f in glob.glob(os.path.join(d, "requests-*.csv")) for r in csv.DictReader(open(f))]
    end = max(float(r["ts"]) for r in reqs)
    return d, st, reqs, end


def check(stage, name):
    for c in stage["checks"]:
        if c["check"] == name:
            return c["value"]
    return None


def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border = HDR, HDR_FILL, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 45


def body(ws, r, values, fmts):
    for i, (v, f) in enumerate(zip(values, fmts), 1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.font, cell.border, cell.number_format = NORM, BOX, f
        cell.alignment = Alignment(horizontal="center")


def verdict_format(ws, col, first, last):
    rng = f"{col}{first}:{col}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{first}="PASS"'], fill=GREEN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{first}="FAIL"'], fill=RED))


wb = Workbook()

# ---------------- Summary ----------------
ws = wb.active
ws.title = "Summary"
ws["A1"] = "Moodle capacity test — concurrently active students"
ws["A1"].font = TITLE
ws["A2"] = "Test date 30 Sep 2026 · Locust staged ramp · each stage held 3.5 min, first 90 s excluded as ramp-up"
ws["A2"].font = NOTE

ws["A4"] = "Capacity (highest stage that passed the SLA)"
ws["A4"].font = BOLD
header(ws, 5, ["Workload", "Capacity (students active at once)", "Limited by", "Details sheet"], [44, 22, 46, 18])
rows = [
    ("Normal use: pages, forum, multiple-choice quiz", "'Normal use'", "Host CPU (Moodle PHP ≈ 22 cores at saturation)"),
    ("Coding quiz: CodeRunner answers run in Jobe", "Coding", "Host CPU, shared by Moodle PHP + Jobe + DB"),
    ("Jupyter notebooks (personal server + kernel each)", "Jupyter", "Host RAM (≈ 180 MB per student)"),
]
for i, (name, sheet, lim) in enumerate(rows):
    r = 6 + i
    body(ws, r, [name, None, lim, sheet.strip("'")], ["@", "#,##0", "@", "@"])
    # Verdict is column L on every detail sheet.
    c = ws.cell(row=r, column=2, value=f'=_xlfn.MAXIFS({sheet}!A:A,{sheet}!L:L,"PASS")')
    c.font, c.number_format = Font(name=F, bold=True, color="008000"), "#,##0"
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
    ws.cell(row=r, column=3).alignment = Alignment(horizontal="left")
ws["A9"] = ("The three capacities are not additive: all workloads share the same 28 CPU threads and 62 GB RAM. "
            "Jupyter stopped at 150 because host free RAM reached the 6 GB safety floor; practical limit ≈ 140.")
ws["A9"].font = NOTE

ws["A11"] = "SLA — a stage PASSES only if every limit holds (edit the blue cells to re-judge the stages)"
ws["A11"].font = BOLD
header(ws, 12, ["Check", "Limit", "Unit", ""], [44, 22, 46, 18])
sla = [("Page loads p95 below", 2.5, "seconds"), ("Graded submission p95 below", 5, "seconds"),
       ("Jupyter server start p95 below", 10, "seconds"), ("Jupyter cell run p95 below", 5, "seconds"),
       ("Error rate below", 0.01, "fraction of requests"), ("Host CPU average below", 0.9, "fraction of 28 threads")]
for i, (n, v, u) in enumerate(sla):
    r = 13 + i
    body(ws, r, [n, v, u], ["@", "0.0%" if v < 1 else "0.0", "@"])
    ws.cell(row=r, column=2).font = INPUT
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
    ws.cell(row=r, column=3).alignment = Alignment(horizontal="left")
ws["B13"].comment = Comment("Limits chosen before the test (user brief: page p95 < 2.5 s, graded submit p95 < 5 s, "
                            "errors < 1 %, host not CPU-pegged).", "Claude")
SLA = {"page": "Summary!$B$13", "submit": "Summary!$B$14", "spawn": "Summary!$B$15", "cell": "Summary!$B$16",
       "err": "Summary!$B$17", "cpu": "Summary!$B$18"}

ws["A20"] = "Hardware condition"
ws["A20"].font = BOLD
hw = [("CPU", "Intel i7-14700, 20 cores / 28 threads"),
      ("RAM", "62 GB total, ≈ 32 GB available to the test (rest used by other services)"),
      ("Disk", "NVMe SSD"),
      ("Host", "Shared with ASR / TTS / diarization / Open WebUI containers (idle during test); no swap activity"),
      ("Web tier", "Apache event MPM + PHP-FPM (112 children), MaxRequestWorkers 1024, KeepAliveTimeout 2 s"),
      ("Database", "MariaDB 11, max_connections 1100, 2 GB buffer pool; file sessions; no Redis/memcached"),
      ("Load generator", "Locust 2.46, 8 processes, on the same host (≈ 1 core at peak)"),
      ("Before tuning", "Apache prefork + mod_php, 256 workers: normal-use capacity was 750 (worker-limited, CPU 10 %)")]
for i, (k, v) in enumerate(hw):
    ws.cell(row=21 + i, column=1, value=k).font = BOLD
    ws.cell(row=21 + i, column=2, value=v).font = NORM

# ---------------- Normal use ----------------
ws = wb.create_sheet("Normal use")
ws["A1"] = "Track 1 — normal use (dashboard, course, reading page, forum, multiple-choice quiz; 5–15 s think time)"
ws["A1"].font = BOLD
cols = ["Students online", "Requests / s", "Page load p95 (s)", "Quiz submit p95 (s)", "Requests (measured window)",
        "Errors", "Error rate", "Host CPU avg", "Moodle CPU (cores)", "DB CPU (cores)", "Free host RAM min (GB)",
        "Verdict"]
header(ws, 3, cols, [12, 11, 12, 12, 14, 9, 10, 11, 12, 11, 13, 10])
data = report("browse")
for i, s in enumerate(data):
    r = 4 + i
    h = s["host"]
    body(ws, r, [s["users"], s["rps"], check(s, "page loads p95") / 1000, check(s, "graded submit p95") / 1000,
                 s["requests"], s["errors"], f"=IF(E{r}=0,0,F{r}/E{r})", h["host_cpu_avg"] / 100,
                 h["moodle_app_cpu_avg"] / 100, h["moodle_db_cpu_avg"] / 100, h["mem_avail_min_gb"],
                 f'=IF(AND(C{r}<{SLA["page"]},D{r}<{SLA["submit"]},G{r}<{SLA["err"]},H{r}<{SLA["cpu"]}),"PASS","FAIL")'],
         ["#,##0", "0.0", "0.00", "0.00", "#,##0", "#,##0", "0.00%", "0.0%", "0.0", "0.0", "0.0", "@"])
last = 3 + len(data)
verdict_format(ws, "L", 4, last)
ws.cell(row=last + 2, column=1, value=f"Source: locust/results/{RUNS['browse']} (stages_report.json, monitor.csv). "
        "CPU columns: averages over the measured window; 1 core = 100 % of one thread.").font = NOTE
ws.freeze_panes = "B4"

# ---------------- Coding ----------------
ws = wb.create_sheet("Coding")
ws["A1"] = "Track 2 — CodeRunner coding quiz (write code 20–40 s, submit, answer run in Jobe sandbox; 10–30 s pause)"
ws["A1"].font = BOLD
cols = ["Students coding", "Requests / s", "Jobe runs / s", "Page load p95 (s)", "Submit + grade p95 (s)",
        "Requests (measured window)", "Errors", "Error rate", "Host CPU avg", "Moodle CPU (cores)",
        "Jobe CPU max (cores)", "Verdict", "DB CPU (cores)"]
header(ws, 3, cols, [12, 11, 11, 12, 13, 14, 9, 10, 11, 12, 12, 10, 11])
data = report("coding")
_, st, reqs, end = stage_windows("coding")
for i, s in enumerate(data):
    r = 4 + i
    h = s["host"]
    t0 = st[i][0] + WARMUP
    t1 = st[i + 1][0] if i + 1 < len(st) else end
    subs = sum(1 for q in reqs if q["name"] == "coding quiz: submit & grade" and q["ok"] == "1" and t0 <= float(q["ts"]) < t1)
    body(ws, r, [s["users"], s["rps"], round(subs / s["window_s"], 1), check(s, "page loads p95") / 1000,
                 check(s, "graded submit p95") / 1000, s["requests"], s["errors"], f"=IF(F{r}=0,0,G{r}/F{r})",
                 h["host_cpu_avg"] / 100, h["moodle_app_cpu_avg"] / 100, h["jobe_cpu_max"] / 100,
                 f'=IF(AND(D{r}<{SLA["page"]},E{r}<{SLA["submit"]},H{r}<{SLA["err"]},I{r}<{SLA["cpu"]}),"PASS","FAIL")',
                 h["moodle_db_cpu_avg"] / 100],
         ["#,##0", "0.0", "0.0", "0.00", "0.00", "#,##0", "#,##0", "0.00%", "0.0%", "0.0", "0.0", "@", "0.0"])
last = 3 + len(data)
verdict_format(ws, "L", 4, last)
ws.cell(row=last + 2, column=1, value=f"Source: locust/results/{RUNS['coding']}. 32,091 attempts graded, average mark "
        "0.75 (3 of 4 sample answers correct) — confirms the code really ran. Jobe's own limit is 16 programs at "
        "once.").font = NOTE
ws.freeze_panes = "B4"

# ---------------- Jupyter ----------------
ws = wb.create_sheet("Jupyter")
ws["A1"] = "Track 3 — Jupyter notebooks (open activity → personal server starts → pandas cell; a cell every 15–30 s)"
ws["A1"].font = BOLD
cols = ["Students with a notebook open", "Server start p50 (s)", "Server start p95 (s)", "Cell run p95 (s)",
        "Requests (measured window)", "Errors", "Error rate", "JupyterHub RAM (GB)", "RAM per student (MB)",
        "Free host RAM min (GB)", "Host CPU avg", "Verdict"]
header(ws, 3, cols, [14, 11, 11, 10, 14, 9, 10, 12, 12, 13, 11, 10])
data = report("jupyter")
d, st, reqs, end = stage_windows("jupyter")
mon = list(csv.DictReader(open(os.path.join(d, "monitor.csv"))))


def pct(v, q):
    v = sorted(v)
    return v[min(len(v) - 1, int(round(q * (len(v) - 1))))]


for i, s in enumerate(data):
    r = 4 + i
    h = s["host"]
    t0, t1 = st[i][0], (st[i + 1][0] if i + 1 < len(st) else end)
    sp = [float(q["ms"]) / 1000 for q in reqs if q["name"].startswith("jupyter: open activity") and t0 <= float(q["ts"]) < t1]
    hub = max(float(x["moodle_jupyterhub_mib"]) for x in mon if t0 <= float(x["ts"]) < t1) / 1024
    body(ws, r, [s["users"], round(pct(sp, .5), 2), round(pct(sp, .95), 2), check(s, "jupyter cell p95") / 1000,
                 s["requests"], s["errors"], f"=IF(E{r}=0,0,F{r}/E{r})", round(hub, 2), f"=H{r}*1024/A{r}",
                 h["mem_avail_min_gb"], h["host_cpu_avg"] / 100,
                 f'=IF(AND(C{r}<{SLA["spawn"]},D{r}<{SLA["cell"]},G{r}<{SLA["err"]},K{r}<{SLA["cpu"]}),"PASS","FAIL")'],
         ["#,##0", "0.0", "0.0", "0.00", "#,##0", "#,##0", "0.00%", "0.0", "#,##0", "0.0", "0.0%", "@"])
last = 3 + len(data)
verdict_format(ws, "L", 4, last)
ws.cell(row=last + 2, column=1, value=f"Source: locust/results/{RUNS['jupyter']}. Server start times cover every "
        "server started during the stage (they start during ramp-up). The test stopped at 150 because free host RAM "
        "reached the 6 GB safety floor.").font = NOTE
ws.freeze_panes = "B4"

out = os.path.join(HERE, "capacity_results.xlsx")
wb.save(out)
print(out)
