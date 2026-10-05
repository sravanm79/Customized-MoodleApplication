"""Video load-test report: VIDEO_REPORT.md + video_results.xlsx from the run directories of run_video_suite.sh.

    .venv/bin/python make_video_report.py   (or python3 with openpyxl)
"""
import glob
import json
import os

HERE = os.path.dirname(os.path.abspath(__file__))
VIDEO = json.load(open(os.path.join(HERE, "setup", "video.json")))
RUNS = [  # label, run dir suffix, bitrate kbps, what it represents
    ("As uploaded", "asis", 448, "the file as uploaded (1720x1080, H.264 314 kbps + AAC 128 kbps)"),
    ("720p HD", "hd720", 1500, "a typical 720p HD encode of the same lecture"),
    ("1080p Full HD", "hd1080", 3000, "a typical 1080p Full HD encode of the same lecture"),
]
SECONDS = VIDEO["video_seconds"]


def run_dir(suffix):
    found = sorted(glob.glob(os.path.join(HERE, "results", f"*-video-remote-{suffix}")))
    return found[-1] if found else None


def mb_for(kbps, seconds=SECONDS):
    return kbps * 1000 * seconds / 8 / 1048576  # MB as Moodle shows them (MiB)


rows = []
summary = []
for label, suffix, kbps, desc in RUNS:
    d = run_dir(suffix)
    stages = json.load(open(os.path.join(d, "stages_report.json")))
    browser = json.load(open(os.path.join(d, "browser", "browser_video.json")))
    passed = [s["users"] for s in stages if s["pass"]]
    cap = 0
    for s in stages:
        if not s["pass"]:
            break
        cap = s["users"]
    fail = next((s["users"] for s in stages if not s["pass"]), None)
    for s in stages:
        c = {x["check"]: x["value"] for x in s["checks"]}
        starts = [v for k, v in s["per_request"].items() if k.startswith("video: start")]
        start_p95 = max((v["p95"] for v in starts), default=None)
        chunk = s["per_request"].get("video: chunk", {})
        rows.append({
            "quality": label, "kbps": kbps, "users": s["users"], "verdict": "PASS" if s["pass"] else "FAIL",
            "stalls_per_100": c.get("video stalls /100 chunks"), "stall_s": c.get("video stall s total"),
            "start_p95_ms": start_p95, "chunk_p95_ms": chunk.get("p95"), "errors_pct": c.get("error rate %"),
            "net_avg": s["host"]["net_tx_mbps_avg"], "net_max": s["host"]["net_tx_mbps_max"],
            "cpu": s["host"]["host_cpu_avg"], "fpm": s["host"]["fpm_max"], "app_mib": s["host"]["moodle_app_mib_max"],
        })
    b = browser
    summary.append({
        "quality": label, "kbps": kbps, "desc": desc, "size_mb": round(mb_for(kbps)), "per_hour_mb": round(mb_for(kbps, 3600)),
        "capacity": cap, "fails_at": fail, "dir": os.path.relpath(d, HERE),
        "browser_where": "SERVER-2, over the LAN" if suffix == "hd1080" else "Moodle host, not over the LAN",
        "browser": {
            "viewers": len(b), "res": f"{b[0]['width']}x{b[0]['height']}",
            "page_ms": round(sum(v["page_load_ms"] for v in b) / len(b)),
            "freezes": sum(v["freezes"] for v in b), "dropped": sum(v["dropped"] for v in b),
            "frames": sum(v["frames"] for v in b), "seek_ms": round(sum(v["seek_ms"] for v in b) / len(b)),
        },
    })

# ---- Maximum-viewers run (as uploaded, ramped until it fails; real Chrome at every stage) ----
MAX = run_dir("max")
maxrows = []
if MAX:
    for s in json.load(open(os.path.join(MAX, "stages_report.json"))):
        c = {x["check"]: x["value"] for x in s["checks"]}
        pr = s["per_request"]
        starts = [v["p95"] for k, v in pr.items() if k.startswith("video: start")]
        bfile = os.path.join(MAX, "browser", str(s["users"]), "browser_video.json")
        b = [v for v in json.load(open(bfile)) if "frames_painted" in v] if os.path.exists(bfile) else []
        avg = lambda k: round(sum(v[k] or 0 for v in b) / len(b)) if b else None  # noqa: E731
        maxrows.append({
            "users": s["users"], "verdict": "PASS" if s["pass"] else "FAIL",
            "stalls_per_100": c.get("video stalls /100 chunks"), "stall_s": c.get("video stall s total"),
            "start_p95_ms": max(starts) if starts else None,
            "ttfb_p50": pr.get("video: chunk first byte", {}).get("p50"), "ttfb_p95": pr.get("video: chunk first byte", {}).get("p95"),
            "chunk_p50": pr.get("video: chunk", {}).get("p50"), "chunk_p95": pr.get("video: chunk", {}).get("p95"),
            "page_p95": pr.get("page: video lecture", {}).get("p95"), "errors_pct": c.get("error rate %"),
            "net_avg": s["host"]["net_tx_mbps_avg"], "net_max": s["host"]["net_tx_mbps_max"],
            "cpu": s["host"]["host_cpu_avg"], "fpm": s["host"]["fpm_max"], "apache": s["host"]["apache_busy_max"],
            "b_page": avg("page_load_ms"), "b_freezes": sum(v["freezes"] for v in b) if b else None,
            "b_frozen_s": round(sum(v["freeze_ms"] for v in b) / 1000, 1) if b else None,
            "b_gap_p95": max((v["frame_gap_p95_ms"] or 0) for v in b) if b else None,
            "b_gap_max": max((v["frame_gap_max_ms"] or 0) for v in b) if b else None,
            "b_late": sum(v["late_frames"] for v in b) if b else None,
            "b_late_pct": round(100 * sum(v["late_frames"] for v in b) / max(1, sum(v["frames_painted"] for v in b)), 2) if b else None,
            "b_skipped": sum(v["frames_skipped"] for v in b) if b else None,
            "b_lag": max(v["playback_lag_ms"] for v in b) if b else None, "b_seek": avg("seek_ms"),
            "b_viewers": len(b),
        })
max_cap = 0
for r in maxrows:
    if r["verdict"] != "PASS":
        break
    max_cap = r["users"]
max_fail = next((r["users"] for r in maxrows if r["verdict"] != "PASS"), None)
if maxrows:  # the as-uploaded capacity is what the maximum run measured
    summary[0].update(capacity=max_cap, fails_at=max_fail)

# ---- Markdown ----
dur = f"{int(SECONDS // 60)} min {int(SECONDS % 60)} s"
md = [f"# Video streaming load test: {VIDEO.get('video_name', 'lecture video')}", "",
      "How many students can watch an uploaded lecture video at the same time, by video quality. Load from a second "
      "LAN machine (SERVER-2) with Locust viewers that stream the file in HTTP Range chunks paced at the bitrate, as a "
      "browser player does (10–30 s buffer, half the sittings seek); real Chrome viewers measured what students see "
      "during the busiest stage. 2 min per stage, first 60 s of each stage excluded.", "",
      "## The video", "",
      "| Property | Value |", "|---|---|",
      f"| File | {VIDEO.get('video_name', '')} |",
      f"| Size | {VIDEO['video_bytes'] / 1048576:.1f} MB ({VIDEO['video_bytes']:,} bytes) |",
      f"| Length | {dur} |",
      "| Picture | 1720×1080, 25 fps, H.264 (314 kbps) |", "| Sound | AAC 128 kbps |",
      "| Total bitrate | **448 kbps** (≈ 192 MB per viewer per hour) |",
      "| Fast start (index at the front) | Yes: playback starts before the file is downloaded |", "",
      "## Results by quality", "",
      "| Quality | Bitrate | Same 44-min lecture would be | Data per viewer-hour | Concurrent viewers OK | Fails at | Real Chrome viewers (busiest stage) |",
      "|---|---|---|---|---|---|---|"]
for s in summary:
    b = s["browser"]
    md.append(f"| {s['quality']} | {s['kbps'] / 1000:g} Mbps | {s['size_mb']} MB | {s['per_hour_mb']} MB | **{s['capacity']}** | "
              f"{s['fails_at'] or 'not reached'} | {b['res']}, {b['freezes']} freezes, {b['dropped']} dropped frames, "
              f"page {b['page_ms']} ms, seek {b['seek_ms']} ms ({s['browser_where']}) |")
if maxrows:
    md += ["", "## Maximum simultaneous viewers (video as uploaded)", "",
           f"Ramped {maxrows[0]['users']} → {maxrows[-1]['users']} viewers, 2.5 min per stage, until it failed. "
           f"**{max_cap} viewers pass; it fails at {max_fail or 'not reached'}.**", "",
           "Server side (all Locust viewers):", "",
           "| Viewers | Verdict | Freezes /100 chunks | Frozen s total | Video start p95 | Chunk first byte p50 / p95 | 1 MB chunk p50 / p95 | Lecture page p95 | Errors | Network out avg / max | Host CPU | PHP workers (/112) | Apache busy |",
           "|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for r in maxrows:
        md.append(f"| {r['users']} | {r['verdict']} | {r['stalls_per_100']} | {r['stall_s']} | {r['start_p95_ms']} ms | "
                  f"{r['ttfb_p50']} / {r['ttfb_p95']} ms | {r['chunk_p50']} / {r['chunk_p95']} ms | {r['page_p95']} ms | "
                  f"{r['errors_pct']} % | {r['net_avg']} / {r['net_max']} Mbps | {r['cpu']} % | {r['fpm']} | {r['apache']} |")
    md += ["", "What a student sees (5 real Chrome viewers on SERVER-2, over the LAN, during each stage; frame timing from "
           "`requestVideoFrameCallback`, which fires for every frame painted on screen):", "",
           "| Viewers | Page load avg | Freezes (frozen s) | Frame gap p95 / worst | Late frames (> 80 ms) | Frames skipped | Furthest behind real time | Seek avg |",
           "|---|---|---|---|---|---|---|---|"]
    for r in maxrows:
        if r["b_viewers"]:
            md.append(f"| {r['users']} | {r['b_page']} ms | {r['b_freezes']} ({r['b_frozen_s']} s) | {r['b_gap_p95']} / "
                      f"{r['b_gap_max']} ms | {r['b_late']} ({r['b_late_pct']} %) | {r['b_skipped']} | {r['b_lag']} ms | {r['b_seek']} ms |")
    md += ["", "How to read the frame columns: the lecture is 25 fps, and headless Chrome paints at about 30 Hz, so "
           "frames normally arrive 33 or 67 ms apart (idle baseline: p95 67 ms, worst 83 ms, behind real time < 50 ms). "
           "A *late frame* (> 80 ms gap) is a visible stutter; *behind real time* is how much playback fell behind the "
           "clock in 60 s of watching (freezes and stutters added up). *Chunk first byte* is how long the server takes "
           "to start answering a video request (queueing in Apache/PHP + Moodle's checks), before any data flows."]
md += ["", "## Every stage", "",
       "| Quality | Viewers | Verdict | Freezes /100 chunks | Video start p95 | Chunk p95 | Errors | Network out avg / max | Host CPU | PHP workers (/112) |",
       "|---|---|---|---|---|---|---|---|---|---|"]
for r in rows:
    md.append(f"| {r['quality']} | {r['users']} | {r['verdict']} | {r['stalls_per_100']} | {r['start_p95_ms'] or '–'} ms | "
              f"{r['chunk_p95_ms'] or '–'} ms | {r['errors_pct']} % | {r['net_avg']} / {r['net_max']} Mbps | {r['cpu']} % | {r['fpm']} |")
md += ["", "Pass = video start p95 < 3 s, < 1 freeze per 100 chunks, errors < 1 %, host CPU < 90 %, no swapping.", "",
       "## What this means", "",
       "- **The limit is the network, not the server.** Every failure happened when the 1 Gbps link was full "
       "(976 Mbps out); host CPU stayed under 11 %, memory and the database were idle, and there were no errors. "
       "Usable capacity ≈ **800 Mbps ÷ bitrate**: as uploaded (448 kbps) **measured: 1,750 pass, 2,000 fail** (see Maximum simultaneous "
       "viewers); 1.5 Mbps ≈ 500; 3 Mbps ≈ 250.",
       "- **Quality (bitrate) decides how many can watch at once.** Doubling the bitrate halves the viewers. This lecture "
       "at 448 kbps already looks sharp at 1720×1080 (slides change slowly), so re-encoding it at a higher bitrate would "
       "cost capacity without a visible gain.",
       "- **Size and length** do not change how many can watch at once; they set storage and total data: "
       "size = bitrate × length (this 44-min lecture: " + ", ".join(f"~{x['size_mb']} MB at {x['kbps'] / 1000:g} Mbps" for x in summary) + "), and every "
       "student watching the whole lecture downloads its full size. Long lectures mean more seeks and resumes, which "
       "the test included (half the sittings seek).",
       "- **When the link is full, everyone suffers:** viewers freeze, start-up takes 4–8 s, and even ordinary pages "
       "slow down (the lecture page took 6.7 s for Chrome viewers on SERVER-2 at 400 × 3 Mbps).",
       "- **PHP workers:** Moodle streams files through PHP. Worker processes peaked at 88 of 112 in the first run "
       "and stayed at the 64-process idle pool in the maximum run (2,000 viewers), so PHP never ran out; it would be "
       "the next limit after a faster network link.", "",
       "## Recommendations", "",
       "1. Upload lectures as **H.264 MP4 with fast start, 720p or 1080p at about 0.5–1.5 Mbps** (screen recordings "
       "compress very well). Command: `ffmpeg -i in.mp4 -c:v libx264 -crf 23 -maxrate 1500k -bufsize 3000k "
       "-vf scale=-2:1080 -c:a aac -b:a 96k -movflags +faststart lecture.mp4`.",
       "2. For more simultaneous viewers: a 10 Gbps (or bonded 2×1 Gbps) network link, or host big videos outside "
       "Moodle (e.g. YouTube unlisted / Google Drive embeds).",
       "3. Plan for the worst moment, e.g. a whole batch opening the same lecture before an exam: viewers × bitrate must "
       "stay under ~800 Mbps.", "",
       "## Caveats", "",
       "- The 250- and 500-viewer stages of the *as uploaded* run overlapped another load test on the same machines "
       "(13:40–13:44); they still passed, so the result is conservative, but their network figures include that other "
       "traffic. The 1,000-viewer stage was clean.",
       "- The higher-quality runs simulate bitrate by pacing the same file; real 720p/1080p files of that bitrate behave "
       "the same on the network and server.",
       "- The maximum run was planned up to 3,000 viewers; the load generator (SERVER-2, which also runs other work) "
       "ran low on memory as the 2,250 stage began and stopped itself. Moodle had already failed at 2,000 with the "
       "link full, and more viewers cannot get more than the 1 Gbps link, so higher stages would only fail harder.",
       "- In the first two runs the real Chrome viewers ran on the Moodle host itself, so they did not feel the network "
       "limit; in the 1080p run they ran on SERVER-2.", "",
       "Raw data: " + ", ".join(f"`{s['dir']}`" for s in summary) + "."]
open(os.path.join(HERE, "VIDEO_REPORT.md"), "w").write("\n".join(md) + "\n")

# ---- Excel ----
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook()
ws = wb.active
ws.title = "Summary"
bold = Font(bold=True)
head = PatternFill("solid", fgColor="0B2C5F")
def header(sheet, cols):
    sheet.append(cols)
    for c in sheet[sheet.max_row]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head; c.alignment = Alignment(wrap_text=True, vertical="top")
ws.append(["Video streaming load test", VIDEO.get("video_name", "")]); ws["A1"].font = Font(bold=True, size=14)
ws.append(["Size (MB)", round(VIDEO["video_bytes"] / 1048576, 1)]); ws.append(["Length", dur])
ws.append(["Resolution", "1720x1080 @ 25 fps"]); ws.append(["Bitrate (kbps)", 448]); ws.append([])
header(ws, ["Quality", "Bitrate (Mbps)", "Same lecture size (MB)", "MB per viewer-hour", "Concurrent viewers OK",
            "Fails at", "Chrome viewers: resolution", "freezes", "dropped frames", "page load ms", "seek ms", "Chrome ran on"])
for s in summary:
    b = s["browser"]
    ws.append([s["quality"], s["kbps"] / 1000, s["size_mb"], s["per_hour_mb"], s["capacity"], s["fails_at"] or "not reached",
               b["res"], b["freezes"], b["dropped"], b["page_ms"], b["seek_ms"], s["browser_where"]])
st = wb.create_sheet("Stages")
header(st, ["Quality", "Bitrate kbps", "Viewers", "Verdict", "Freezes /100 chunks", "Stall s total", "Video start p95 ms",
            "Chunk p95 ms", "Error %", "Net out avg Mbps", "Net out max Mbps", "Host CPU %", "PHP workers max", "moodle_app MiB"])
red = PatternFill("solid", fgColor="F8D7DA"); green = PatternFill("solid", fgColor="D1E7DD")
for r in rows:
    st.append([r["quality"], r["kbps"], r["users"], r["verdict"], r["stalls_per_100"], r["stall_s"], r["start_p95_ms"],
               r["chunk_p95_ms"], r["errors_pct"], r["net_avg"], r["net_max"], r["cpu"], r["fpm"], r["app_mib"]])
    st.cell(st.max_row, 4).fill = green if r["verdict"] == "PASS" else red
sheets = [ws, st]
if maxrows:
    mx = wb.create_sheet("Max viewers")
    sheets.append(mx)
    keys = list(maxrows[0])
    header(mx, keys)
    for r in maxrows:
        mx.append([r[k] for k in keys])
        mx.cell(mx.max_row, 2).fill = green if r["verdict"] == "PASS" else red
for sheet in sheets:
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 18
wb.save(os.path.join(HERE, "video_results.xlsx"))
print("wrote VIDEO_REPORT.md and video_results.xlsx")
